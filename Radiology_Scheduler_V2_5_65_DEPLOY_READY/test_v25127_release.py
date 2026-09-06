from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parent
APP = (ROOT / 'app.py').read_text(encoding='utf-8')
ENG = (ROOT / 'scheduler_engine.py').read_text(encoding='utf-8')
SQL = (ROOT / 'SUPABASE_MIGRATION_V2_5_125_LT_UX_LIFECYCLE_DREAMTEAM.sql').read_text(encoding='utf-8')
GUIDE = (ROOT / 'SENIOR_USABILITY_GUIDE_LT.md').read_text(encoding='utf-8')
README = (ROOT / 'README_LT.md').read_text(encoding='utf-8')
UPDATE = (ROOT / 'ATNAUJINIMAS_V2_5_127.md').read_text(encoding='utf-8')


def _manual_text():
    doc = Document(ROOT / 'SHIFT_HAPPENS_SENIUNES_VADOVAS_V2_5_127.docx')
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(' | '.join(cell.text for cell in row.cells))
    return '\n'.join(chunks)


def test_release_version_and_lithuanian_ui():
    assert 'APP_VERSION = "2.5.127 LIETUVIŠKA UX + PRISTATYMO VADOVAS"' in APP
    assert 'lang = "LT"' in APP
    assert 'Kalba / Language' not in APP
    assert 'st.sidebar.caption("Kalba: lietuvių")' in APP


def test_resident_copy_is_plain_language():
    assert '"hard_unavailable":"Dirbti negaliu"' in APP
    assert '"soft_free":"Noriu laisvos"' in APP
    assert '"preferred":"Pageidauju dirbti"' in APP
    assert 'Pateikimo vieta skaičiuojama iš naujo kiekvienam grafikui ir galioja būtent tam grafikui' in APP
    assert 'Šio mėnesio pateikimo vieta uždirba prioritetą tik kitam grafikui' not in APP


def test_public_hierarchy_contains_no_hidden_planning_layer():
    public = '\n'.join([GUIDE, README, UPDATE, _manual_text()]).lower()
    for forbidden in ('privat', 'dream team', 'komandos tiksl'):
        assert forbidden not in public
    assert 'saugumas, įmanomumas ir privalomas padengimas' in public
    assert 'maksimalus visų rezidentų pageidavimų išpildymas' in public
    assert 'pateikimo eilė' in public


def test_goal_is_100_with_natural_complex_month_example():
    public = '\n'.join([GUIDE, README, UPDATE, _manual_text()])
    assert '100 %' in public
    assert '93 %' in public
    assert '90–95' not in public
    assert '90-95' not in public
    assert 'nėra nustatyta riba' not in public.lower()
    assert 'neturi 90' not in public.lower()


def test_research_comparator_is_migrated():
    assert 'names.append("AVAILABLE GPT + HUMAN vs MY ENGINE")' not in APP
    assert 'with st.expander("Grafikų sudarymo metodų palyginimas"' in APP
    assert 'Bendrinis DI + seniūnė' in APP
    assert 'Specializuotas grafiko variklis' in APP


def test_lifecycle_14_15_16():
    assert "when now()<v_pref_close then 'preferences'" in SQL
    assert "when now()<v_swap_open then 'senior_build'" in SQL
    assert "when now()<v_swap_close then 'swaps'" in SQL
    assert 'make_date(py,pm,15)' in SQL
    assert '15 d. 00:00–16 d. 00:00' in README or '15 d. 00:00 – 16 d. 00:00' in GUIDE
    assert 'Nuo 16 d. 00:00' in GUIDE


def test_submission_rank_only_after_maximum_public_result():
    assert 'Sistema visada pirmiausia siekia įvykdyti kuo daugiau visų rezidentų pageidavimų.' in APP
    assert 'Sistema visada pirmiausia siekia 100% bendro pageidavimų išpildymo.' in APP
    assert 'maximum total fulfilment are locked first' in ENG or 'total fulfilment lock' in ENG
    assert 'submission priority lock' in ENG


def test_private_refinement_cannot_buy_public_wish_loss():
    # Private variables are neutral during the ordinary public solve.
    assert 'cost=-1000000.0' not in ENG
    assert 'ZERO weight during the group solve' in ENG
    # Fast phase freezes every exact visible resident request before private refinement.
    assert 'Freeze exact visible request outcomes for all residents' in ENG
    assert 'No visible resident request can be' in ENG
    # Legacy path also defers private refinement until after public SOFT + post stages.
    assert ENG.index('Final private-only refinement for the legacy single-pass solver') > ENG.index('# Stage I')
    assert 'private refine public shift lock' in ENG
    assert 'private refine exposure lock' in ENG
    # The old early private optimum lock must not exist.
    assert 'priority coloc optimum lock' not in ENG


def test_private_controls_remain_out_of_simple_mode():
    assert 'if active_user==SENIOR_INITIALS and preference_target==SENIOR_INITIALS and advanced_mode:' in APP
    assert 'render_sp_dream_team_settings_v25125(year,month)' in APP
    assert 'render_sp_private_pair_preferences(year,month)' in APP


def test_monthly_private_targets_are_not_weekly_rules():
    assert 'dream_team_centro_count' in ENG
    assert 'dream_team_centro_target' in ENG
    assert 'monthly occurrence target' in ENG
    assert 'priority coloc max one block week' not in ENG
